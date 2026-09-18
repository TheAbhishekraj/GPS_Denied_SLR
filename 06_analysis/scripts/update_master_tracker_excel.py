#!/usr/bin/env python3
"""
update_master_tracker_excel.py — Build PROJECT_MASTER_TRACKER.xlsx
"""

import os, csv, glob, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = r"E:\GPS_Denied_SLR"
EXCEL_PATH = os.path.join(REPO, "08_docs", "PROJECT_MASTER_TRACKER.xlsx")

HEADER_FILL = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9")
)
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
def build_inventory_rows():
    key_files = [
        ("00_scope/PROTOCOL.md", "00_scope", "Markdown", "Core Protocol", "Primary SLR protocol RQs & criteria.", "Verified"),
        ("00_scope/REGISTRATION.md", "00_scope", "Markdown", "Core Protocol", "PRISMA-P compliant registration.", "Verified"),
        ("00_scope/research_protocol.md", "00_scope", "Markdown", "Core Protocol", "Research design & taxonomies.", "Verified"),
        ("00_scope/screening_criteria_v2.md", "00_scope", "Markdown", "Core Protocol", "Phase 2 screening rules.", "Verified"),
        ("00_scope/quality_appraisal_rubric.md", "00_scope", "Markdown", "Core Protocol", "7-parameter quality rubric.", "Verified"),
        ("00_scope/SEARCH_STRINGS.md", "00_scope", "Markdown", "Core Protocol", "Search queries for IEEE/Scopus.", "Verified"),
        ("00_scope/human_validation_protocol.md", "00_scope", "Markdown", "Core Protocol", "SOP for 34-paper human validation.", "Verified"),
        ("01_data_raw/ieee_xplore.csv", "01_data_raw", "CSV", "Raw Data", "Raw search export from IEEE Xplore.", "Verified"),
        ("01_data_raw/scopus.csv", "01_data_raw", "CSV", "Raw Data", "Raw search export from Scopus.", "Verified"),
        ("01_data_raw/SEARCH_LOG.md", "01_data_raw", "Markdown", "Audit Log", "Database search execution logs.", "Verified"),
        ("02_data_processed/deduplicated_master.csv", "02_data_processed", "CSV", "Processed Data", "Deduplicated master records.", "Verified"),
        ("02_data_processed/screened_included_v2.csv", "02_data_processed", "CSV", "Core Corpus", "636 screened included papers.", "Verified"),
        ("02_data_processed/extracted_master_v2.csv", "02_data_processed", "CSV", "Core Corpus", "171 core extracted papers with QA.", "Verified"),
        ("02_data_processed/MASTER_EVIDENCE_V1.csv", "02_data_processed", "CSV", "Master Table", "63-column evidence matrix.", "Active Target"),
        ("05_papers_fulltext/", "05_papers_fulltext", "PDF Directory", "Fulltext Corpus", "280 full-text PDFs (REC_XXXX.pdf).", "280 PDFs"),
        ("06_analysis/scripts/01_deduplicate.py", "06_analysis", "Python Script", "Pipeline Executable", "Deduplication algorithm.", "Verified"),
        ("06_analysis/scripts/09_build_evidence_workbooks.py", "06_analysis", "Python Script", "Pipeline Executable", "Generates evidence workbooks.", "Verified"),
        ("06_analysis/scripts/10_commit_extraction.py", "06_analysis", "Python Script", "Pipeline Executable", "Extraction validator and committer.", "Verified"),
        ("06_analysis/scripts/16_execute_phase7_extraction.py", "06_analysis", "Python Script", "Pipeline Executable", "PDF parser and QA scoring runner.", "Verified"),
        ("06_analysis/scripts/18_integrate_batch_v2.py", "06_analysis", "Python Script", "Pipeline Executable", "Batch staging integrator.", "Verified"),
        ("07_manuscript/GPS_Denied_SLR_Manuscript_V1_171.tex", "07_manuscript", "LaTeX Source", "Manuscript Asset", "Primary IEEE Transactions TeX file.", "Verified"),
        ("07_manuscript/GPS_Denied_SLR_Manuscript_V1_171.pdf", "07_manuscript", "PDF Document", "Manuscript Asset", "Compiled manuscript PDF.", "Verified"),
        ("07_manuscript/BUILD.md", "07_manuscript", "Markdown", "Build Spec", "LaTeX compilation guide.", "Verified"),
        ("08_docs/PHASE1_MASTER.csv", "08_docs", "CSV", "Master Tracker", "339-row tracking CSV.", "Verified"),
        ("08_docs/priority_300_download.csv", "08_docs", "CSV", "Master Tracker", "300 priority download tracker.", "241 Y / 59 N"),
        ("08_docs/remaining_59_downloads.csv", "08_docs", "CSV", "Master Tracker", "Remaining 59 pending downloads.", "59 Pending"),
        ("08_docs/RECONCILIATION_AUDIT.csv", "08_docs", "CSV", "Audit Report", "300-row audit classification.", "Verified"),
        ("08_docs/TO_DOWNLOAD_MANUAL.csv", "08_docs", "CSV", "Audit Report", "Filtered manual download target.", "59 Rows"),
        ("08_docs/ORPHAN_PDFS.txt", "08_docs", "Text File", "Audit Report", "List of 39 orphan PDFs.", "39 Rows"),
        ("08_docs/extraction_validation_sample_v1.csv", "08_docs", "CSV", "Sample Set", "34-paper verification sample.", "Verified"),
        ("README.md", "ROOT", "Markdown", "Documentation", "Repository main guide.", "Verified"),
        ("RULINGS.md", "ROOT", "Markdown", "Core Specification", "Methodological rulings.", "Verified"),
        ("CHANGELOG.md", "ROOT", "Markdown", "Audit Log", "Repository changelog.", "Verified")
    ]
    rows = []
    for idx, (rel_p, folder, ftype, sig, summary, auto_st) in enumerate(key_files, start=1):
        abs_p = os.path.abspath(os.path.join(REPO, rel_p))
        rows.append({
            "file_id": f"FILE_{idx:04d}", "relative_path": rel_p, "absolute_path": abs_p,
            "folder_category": folder, "file_type": ftype, "significance_level": sig,
            "content_summary": summary, "auto_status": auto_st,
            "manual_verification_status": "[ ] Unverified", "manual_remarks": ""
        })
    return rows


def build_corpus_rows():
    phase1_master = os.path.join(REPO, "08_docs", "PHASE1_MASTER.csv")
    rows = []
    if os.path.exists(phase1_master):
        with open(phase1_master, encoding="utf-8") as f:
            for r in csv.DictReader(f):
                pdf_fname = r.get("pdf_filename", "NONE")
                abs_pdf = os.path.join(REPO, "05_papers_fulltext", pdf_fname) if pdf_fname != "NONE" else "N/A (Missing)"
                rows.append({
                    "id": r.get("id", ""), "priority_rank": r.get("priority_rank", ""),
                    "in_priority_300": r.get("in_priority_300", ""), "done_Y_N": r.get("done_Y_N", ""),
                    "pdf_present": r.get("pdf_present", ""), "in_extracted_master": r.get("in_extracted_master", ""),
                    "extraction_pending": r.get("extraction_pending", ""), "status": r.get("status", ""),
                    "source": r.get("source", ""), "year": r.get("year", ""), "title": r.get("title", ""),
                    "doi": r.get("doi", ""), "venue": r.get("venue", ""), "pdf_filename": pdf_fname,
                    "pdf_absolute_path": abs_pdf, "manual_pdf_verification": "[ ] Unverified", "manual_remarks": ""
                })
    return rows

def build_sample_34_rows():
    sample_path = os.path.join(REPO, "08_docs", "MASTER_EVIDENCE_VERIFY_SAMPLE.csv")
    master_path = os.path.join(REPO, "02_data_processed", "extracted_master_v2.csv")
    master_dict = {}
    if os.path.exists(master_path):
        with open(master_path, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f): master_dict[r["id"].strip()] = r

    rows = []
    if os.path.exists(sample_path):
        with open(sample_path, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                rid = r.get("id", "").strip()
                m = master_dict.get(rid, {})
                pdf_p = os.path.join(REPO, "05_papers_fulltext", f"{rid}.pdf")
                wb_p = os.path.join(REPO, "06_analysis", "output", "workbooks", f"{rid}.md")
                sum_p = os.path.join(REPO, "03_extraction", "per_paper", f"{rid}.md")
                rows.append({
                    "id": rid, "stratum": r.get("approach_family", "") or "Core Extraction",
                    "title": r.get("title", "") or m.get("title", ""),
                    "year": r.get("year", "") or m.get("year", ""), "venue": r.get("venue", "") or m.get("venue", ""),
                    "qa_tier": m.get("qa_tier", "N/A"), "qa_total": m.get("qa_total", "N/A"),
                    "pdf_absolute_path": pdf_p, "workbook_absolute_path": wb_p, "summary_md_path": sum_p,
                    "check_title_doi": "[ ] Pending", "check_pdf_readable": "[ ] Pending",
                    "check_taxonomy": "[ ] Pending", "check_qa_scores": "[ ] Pending",
                    "check_numeric_metrics": "[ ] Pending",
                    "human_verdict": r.get("human_verdict") or "[ ] Unverified",
                    "auditor_remarks": r.get("human_notes") or ""
                })
    return rows

def build_folder_guide_rows():
    folders = [
        ("00_scope", 12, "Phase 0 - Protocol", "Markdown, Text", "Defines RQs, search queries, screening rules, QA rubric."),
        ("01_data_raw", 6, "Phase 1 - Retrieval", "CSV, Markdown", "Raw database exports from IEEE Xplore and Scopus."),
        ("02_data_processed", 25, "Phase 2 - Deduplication & Screening", "CSV, Markdown", "Deduplicated master, screened included (636), extracted master (171)."),
        ("03_extraction", 1, "Phase 7 - Data Extraction", "Markdown Summaries", "Per-paper structured Markdown summaries."),
        ("05_papers_fulltext", 280, "Phase 1 - Full-Text Corpus", "PDF Documents", "Canonical full-text PDF repository (REC_XXXX.pdf)."),
        ("06_analysis", 296, "Phase 7/8 - Analysis & Pipeline", "Python Scripts, Workbooks", "Automated analysis scripts, staging batch integrators, 171 evidence workbooks."),
        ("07_manuscript", 21, "Phase 13 - Manuscript Build", "LaTeX, PDF, Markdown", "IEEE Transactions LaTeX manuscript source, compiled PDF, figures, build specs."),
        ("08_docs", 744, "All Phases - Documentation & Tracking", "CSV, Excel, Markdown", "Project master Excel tracker, priority 300 download sheets, audit logs."),
        ("ROOT", 35, "Project Governance", "Markdown, Configuration", "README, RULINGS.md, CHANGELOG.md, repository specifications.")
    ]
    return [{"folder_name": n, "file_count": c, "slr_phase": p, "primary_formats": f, "purpose_and_significance": d} for n, c, p, f, d in folders]
def main():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("Project_File_Inventory")
    ws1.append(["File ID", "Relative Path", "Absolute Path", "Folder Category", "File Format", "Significance Level", "Content Description & SLR Role", "Auto Status", "Manual Verification Status", "Manual Audit Remarks"])
    for r in build_inventory_rows():
        ws1.append([r["file_id"], r["relative_path"], r["absolute_path"], r["folder_category"], r["file_type"], r["significance_level"], r["content_summary"], r["auto_status"], r["manual_verification_status"], r["manual_remarks"]])

    ws2 = wb.create_sheet("Phase1_Master_Corpus")
    ws2.append(["REC ID", "Priority Rank", "Priority 300?", "Download Done?", "PDF Present?", "In Extracted Master?", "Extraction Pending?", "Reconciliation Status", "Database Source", "Year", "Paper Title", "DOI", "Venue", "PDF Filename", "PDF Absolute Path", "Manual PDF Verification", "Human Audit Remarks"])
    for r in build_corpus_rows():
        ws2.append([r["id"], r["priority_rank"], r["in_priority_300"], r["done_Y_N"], r["pdf_present"], r["in_extracted_master"], r["extraction_pending"], r["status"], r["source"], r["year"], r["title"], r["doi"], r["venue"], r["pdf_filename"], r["pdf_absolute_path"], r["manual_pdf_verification"], r["manual_remarks"]])

    ws3 = wb.create_sheet("Verification_Sample_34")
    ws3.append(["REC ID", "Stratum", "Paper Title", "Year", "Venue", "QA Tier", "QA Score", "PDF Absolute Path", "Workbook Absolute Path", "Summary MD Path", "Check: Title/DOI Match", "Check: PDF Readable", "Check: Taxonomy Acc.", "Check: QA Scores", "Check: Metrics Verified", "Human Audit Verdict", "Auditor Remarks"])
    for r in build_sample_34_rows():
        ws3.append([r["id"], r["stratum"], r["title"], r["year"], r["venue"], r["qa_tier"], r["qa_total"], r["pdf_absolute_path"], r["workbook_absolute_path"], r["summary_md_path"], r["check_title_doi"], r["check_pdf_readable"], r["check_taxonomy"], r["check_qa_scores"], r["check_numeric_metrics"], r["human_verdict"], r["auditor_remarks"]])

    ws4 = wb.create_sheet("Directory_Structure_Guide")
    ws4.append(["Folder Name", "File Count", "SLR Phase", "Primary Formats", "Purpose & Significance"])
    for r in build_folder_guide_rows():
        ws4.append([r["folder_name"], r["file_count"], r["slr_phase"], r["primary_formats"], r["purpose_and_significance"]])

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        for col_idx in range(1, ws.max_column + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = HEADER_FILL; c.font = HEADER_FONT
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_idx in range(2, ws.max_row + 1):
            fill = ZEBRA_FILL if row_idx % 2 == 0 else None
            for col_idx in range(1, ws.max_column + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                if fill: c.fill = fill
                c.border = BORDER_THIN; c.font = Font(name="Calibri", size=10)
                v = str(c.value or "")
                if v in ("Y", "N", "[ ] Unverified", "[ ] Pending", "N/A", "NONE") or len(v) < 8:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")

        for col in ws.columns:
            max_len = max(len(str(c.value or "").split("\n")[0]) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    for ws in wb.worksheets: ws.row_dimensions[1].height = 28
    wb.save(EXCEL_PATH)
    print(f"Successfully generated master tracker Excel at: {EXCEL_PATH}")

if __name__ == "__main__":
    main()

